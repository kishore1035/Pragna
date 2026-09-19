"""Tests for the PDF and PPTX edit/read additions to app.document_generator.

Word (.docx) and Excel (.xlsx) already support real in-place editing via
python-docx/openpyxl, which is exercised via the /api/documents/edit route
tests. These tests cover the previously-missing parity for PDF (structure
sidecar based) and PPTX (direct python-pptx editing) create -> edit -> read
round trips.
"""
from pathlib import Path

import pytest

from app import document_generator as dg


# ─────────────────────────────────────────────────────────────────────────
# PPTX: create -> edit -> read
# ─────────────────────────────────────────────────────────────────────────
def test_create_and_read_presentation(tmp_path):
    filepath = tmp_path / "deck.pptx"
    structure = {
        "title": "Quarterly Review",
        "sections": [
            {"heading": "Highlights", "bullets": ["Revenue up 12%", "Two new hires"], "paragraphs": [], "table": None},
        ],
    }
    dg._build_pptx(structure, str(filepath))
    assert filepath.exists()

    result = dg.read_presentation(str(filepath))
    assert result["success"] is True
    titles = [s["title"] for s in result["slides"]]
    assert "Quarterly Review" in titles
    assert "Highlights" in titles
    highlights_slide = next(s for s in result["slides"] if s["title"] == "Highlights")
    assert "Revenue up 12%" in highlights_slide["bullets"]


def test_edit_presentation_add_slide(tmp_path):
    filepath = tmp_path / "deck.pptx"
    dg._build_pptx({"title": "Deck", "sections": []}, str(filepath))

    result = dg.edit_presentation(
        str(filepath), "add_slide", title="New Slide", bullets=["point one", "point two"]
    )
    assert result["success"] is True

    read_back = dg.read_presentation(str(filepath))
    titles = [s["title"] for s in read_back["slides"]]
    assert "New Slide" in titles
    new_slide = next(s for s in read_back["slides"] if s["title"] == "New Slide")
    assert new_slide["bullets"] == ["point one", "point two"]


def test_edit_presentation_update_slide(tmp_path):
    filepath = tmp_path / "deck.pptx"
    structure = {
        "title": "Deck",
        "sections": [{"heading": "Old Title", "bullets": ["old bullet"], "paragraphs": [], "table": None}],
    }
    dg._build_pptx(structure, str(filepath))

    read_before = dg.read_presentation(str(filepath))
    slide_index = next(i for i, s in enumerate(read_before["slides"]) if s["title"] == "Old Title")

    result = dg.edit_presentation(
        str(filepath), "update_slide", slide_index=slide_index, title="Updated Title", bullets=["new bullet"]
    )
    assert result["success"] is True

    read_after = dg.read_presentation(str(filepath))
    titles = [s["title"] for s in read_after["slides"]]
    assert "Updated Title" in titles
    assert "Old Title" not in titles
    updated_slide = next(s for s in read_after["slides"] if s["title"] == "Updated Title")
    assert updated_slide["bullets"] == ["new bullet"]


def test_edit_presentation_replace_text(tmp_path):
    filepath = tmp_path / "deck.pptx"
    structure = {
        "title": "Cats Overview",
        "sections": [{"heading": "About Cats", "bullets": ["Cats are great"], "paragraphs": [], "table": None}],
    }
    dg._build_pptx(structure, str(filepath))

    result = dg.edit_presentation(str(filepath), "replace_text", search="Cats", replace="Dogs")
    assert result["success"] is True

    read_back = dg.read_presentation(str(filepath))
    all_text = " ".join(s["title"] for s in read_back["slides"]) + " " + " ".join(
        b for s in read_back["slides"] for b in s["bullets"]
    )
    assert "Cats" not in all_text
    assert "Dogs" in all_text


def test_edit_presentation_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        dg.edit_presentation(str(tmp_path / "missing.pptx"), "add_slide", title="x")


def test_edit_presentation_unknown_action_raises(tmp_path):
    filepath = tmp_path / "deck.pptx"
    dg._build_pptx({"title": "Deck", "sections": []}, str(filepath))
    with pytest.raises(ValueError):
        dg.edit_presentation(str(filepath), "not_a_real_action")


# ─────────────────────────────────────────────────────────────────────────
# PDF: create (writes sidecar) -> edit -> read
# ─────────────────────────────────────────────────────────────────────────
def test_create_pdf_writes_structure_sidecar(tmp_path):
    filepath = tmp_path / "report.pdf"
    structure = {
        "title": "Annual Report",
        "sections": [{"heading": "Intro", "bullets": ["point a"], "paragraphs": [], "table": None}],
    }
    dg._build_pdf(structure, str(filepath))

    sidecar = Path(str(filepath) + ".structure.json")
    assert sidecar.exists()


def test_edit_pdf_append_section(tmp_path):
    filepath = tmp_path / "report.pdf"
    structure = {
        "title": "Annual Report",
        "sections": [{"heading": "Intro", "bullets": ["point a"], "paragraphs": [], "table": None}],
    }
    dg._build_pdf(structure, str(filepath))

    result = dg.edit_pdf_document(
        str(filepath), "append_section", heading="Conclusion", paragraphs=["We did well."], bullets=["win 1"]
    )
    assert result["success"] is True
    assert filepath.exists()

    read_back = dg.read_pdf_document(str(filepath))
    assert "Conclusion" in read_back["content"]
    assert "We did well." in read_back["content"]


def test_edit_pdf_replace_text(tmp_path):
    filepath = tmp_path / "report.pdf"
    structure = {
        "title": "Cats Report",
        "sections": [{"heading": "About Cats", "bullets": ["Cats are great"], "paragraphs": [], "table": None}],
    }
    dg._build_pdf(structure, str(filepath))

    result = dg.edit_pdf_document(str(filepath), "replace_text", search="Cats", replace="Dogs")
    assert result["success"] is True

    read_back = dg.read_pdf_document(str(filepath))
    assert "Dogs" in read_back["content"]
    assert "Cats" not in read_back["content"]


def test_edit_pdf_without_sidecar_falls_back_to_extracted_text(tmp_path):
    # Simulate a PDF that wasn't generated by this engine (no sidecar present).
    filepath = tmp_path / "external.pdf"
    dg._build_pdf({"title": "External Doc", "sections": [{"heading": "Body", "bullets": ["existing text"], "paragraphs": [], "table": None}]}, str(filepath))
    sidecar = Path(str(filepath) + ".structure.json")
    sidecar.unlink()  # remove sidecar to force fallback path

    result = dg.edit_pdf_document(str(filepath), "append_section", heading="Appendix", paragraphs=["new info"])
    assert result["success"] is True

    read_back = dg.read_pdf_document(str(filepath))
    assert "Appendix" in read_back["content"]
    assert "new info" in read_back["content"]


def test_edit_pdf_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        dg.edit_pdf_document(str(tmp_path / "missing.pdf"), "append_section", heading="x")


def test_edit_pdf_unknown_action_raises(tmp_path):
    filepath = tmp_path / "report.pdf"
    dg._build_pdf({"title": "Doc", "sections": []}, str(filepath))
    with pytest.raises(ValueError):
        dg.edit_pdf_document(str(filepath), "not_a_real_action")


# ─────────────────────────────────────────────────────────────────────────
# Regression: interleaved paragraphs/bullets/table must render in the order
# they appeared in the source text, and no content type may be silently
# dropped. Previously each builder rendered "all paragraphs, then the
# table, then all bullets" regardless of source order, and the xlsx and
# pptx fallback paths only used ONE of bullets/paragraphs/table per section
# (whichever was checked first), discarding the rest.
# ─────────────────────────────────────────────────────────────────────────
_INTERLEAVED_CONTENT = """# Quarterly Report

## Overview
Intro paragraph.

- First bullet
- Second bullet

Closing paragraph.

| Region | Growth |
|--------|--------|
| NA | 12% |

- Bullet after the table
"""

# No table -- exercises the xlsx "Summary" fallback sheet specifically,
# rather than the per-section-table sheet path.
_INTERLEAVED_CONTENT_NO_TABLE = """# Quarterly Report

## Overview
Intro paragraph.

- First bullet
- Second bullet

Closing paragraph.
"""


def test_parse_markdown_outline_preserves_source_order():
    structure = dg._parse_markdown_outline(_INTERLEAVED_CONTENT, original_prompt="Report")
    section = structure["sections"][0]
    block_summary = [(b["type"], b.get("text") or b.get("rows")) for b in section["blocks"]]
    assert block_summary == [
        ("paragraph", "Intro paragraph."),
        ("bullet", "First bullet"),
        ("bullet", "Second bullet"),
        ("paragraph", "Closing paragraph."),
        ("table", [["Region", "Growth"], ["NA", "12%"]]),
        ("bullet", "Bullet after the table"),
    ]


def test_build_docx_preserves_interleaved_order(tmp_path):
    filepath = tmp_path / "report.docx"
    structure = dg._parse_markdown_outline(_INTERLEAVED_CONTENT, original_prompt="Report")
    dg._build_docx(structure, str(filepath))

    from docx import Document
    from docx.text.paragraph import Paragraph as DocxParagraph

    doc = Document(str(filepath))
    seen = []
    for child in doc.element.body:
        if child.tag.endswith("}p"):
            text = DocxParagraph(child, doc).text.strip()
            if text:
                seen.append(("p", text))
        elif child.tag.endswith("}tbl"):
            seen.append(("table", None))

    # Skip the title/subtitle/heading paragraphs added before section content.
    content_start = next(i for i, (_, t) in enumerate(seen) if t == "Intro paragraph.")
    assert seen[content_start:] == [
        ("p", "Intro paragraph."),
        ("p", "First bullet"),
        ("p", "Second bullet"),
        ("p", "Closing paragraph."),
        ("table", None),
        ("p", "Bullet after the table"),
    ]


def test_build_xlsx_summary_fallback_does_not_drop_paragraphs(tmp_path):
    filepath = tmp_path / "report.xlsx"
    structure = dg._parse_markdown_outline(_INTERLEAVED_CONTENT_NO_TABLE, original_prompt="Report")
    dg._build_xlsx(structure, str(filepath))

    from openpyxl import load_workbook

    wb = load_workbook(str(filepath))
    ws = wb["Summary"]
    contents = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "Intro paragraph." in contents
    assert "Closing paragraph." in contents
    assert "First bullet" in contents


def test_build_pptx_slide_includes_paragraphs_and_bullets(tmp_path):
    filepath = tmp_path / "deck.pptx"
    structure = dg._parse_markdown_outline(_INTERLEAVED_CONTENT, original_prompt="Report")
    dg._build_pptx(structure, str(filepath))

    result = dg.read_presentation(str(filepath))
    slide = next(s for s in result["slides"] if s["title"] == "Overview")
    assert "Intro paragraph." in slide["bullets"]
    assert "First bullet" in slide["bullets"]
    assert "Closing paragraph." in slide["bullets"]
